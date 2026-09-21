"""Raw text parsers for resume ingestion (.pdf, .docx, and raw text).

OWNER: Person A (Dy David)
"""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from PDF file bytes using pdfminer.six."""
    try:
        from pdfminer.high_level import extract_text
        text = extract_text(io.BytesIO(file_bytes))
        return text.strip() if text else ""
    except Exception as err:
        logger.error("Failed to parse PDF file: %s", err)
        raise ValueError(f"Could not extract text from PDF: {err}") from err


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from DOCX file bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs).strip()
    except Exception as err:
        logger.error("Failed to parse DOCX file: %s", err)
        raise ValueError(f"Could not extract text from DOCX: {err}") from err


def parse_resume_bytes(file_bytes: bytes, filename: str) -> str:
    """Route resume file bytes to the appropriate parser based on extension or magic bytes."""
    if not file_bytes:
        raise ValueError("Resume file is empty")

    ext = filename.lower().strip().split(".")[-1] if "." in filename else ""

    raw_text = ""

    # 1. Magic byte check for DOCX (zip file starting with PK)
    if file_bytes.startswith(b"PK\x03\x04") or ext in ("docx", "doc"):
        try:
            raw_text = extract_text_from_docx(file_bytes)
        except Exception as err:
            logger.warning("DOCX parser failed: %s. Trying fallback text decoder.", err)

    # 2. Magic byte check for PDF
    if not raw_text and (file_bytes.startswith(b"%PDF") or ext == "pdf"):
        try:
            raw_text = extract_text_from_pdf(file_bytes)
        except Exception as err:
            logger.warning("PDF parser failed: %s. Trying fallback text decoder.", err)

    # 3. Text decoding fallback
    if not raw_text:
        try:
            raw_text = file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            raw_text = file_bytes.decode("utf-8", errors="ignore").strip()

    if not raw_text or not raw_text.strip():
        raise ValueError("Resume file yielded no extractable text")

    return raw_text

